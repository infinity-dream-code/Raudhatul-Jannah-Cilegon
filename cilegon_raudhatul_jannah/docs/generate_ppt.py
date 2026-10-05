# -*- coding: utf-8 -*-
"""Generate Facepay Cilegon Raudhatul Jannah system documentation PowerPoint."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
# Brand colors (Cilegon Raudhatul Jannah design tokens)
BRAND = RGBColor(0x1F, 0x6B, 0x4A)
BRAND_DEEP = RGBColor(0x13, 0x48, 0x33)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
PAPER = RGBColor(0xF7, 0xF5, 0xF1)
INK = RGBColor(0x14, 0x20, 0x1A)
MUTED = RGBColor(0x5C, 0x6B, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xE6, 0xF2, 0xEC)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=PAPER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return shape


def add_topbar(slide, title_text="Facepay Cilegon Raudhatul Jannah"):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.72))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRAND
    bar.line.fill.background()
    tb = bar.text_frame
    tb.word_wrap = True
    p = tb.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "  " + title_text
    set_run(run, 16, True, WHITE)
    return bar


def add_footer(slide, page, total):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(12), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"Facepay Cilegon Raudhatul Jannah  ·  Dokumentasi Sistem  ·  {page}/{total}"
    set_run(run, 11, False, MUTED)
    p.alignment = PP_ALIGN.RIGHT


def add_title(slide, text, top=Inches(1.0), size=32):
    box = slide.shapes.add_textbox(Inches(0.6), top, Inches(12.2), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size, True, BRAND_DEEP)
    return box


def add_bullets(slide, items, left=Inches(0.7), top=Inches(1.8), width=Inches(12),
                size=16, spacing=True):
    box = slide.shapes.add_textbox(left, top, width, Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        if isinstance(item, tuple):
            level, text = item
            p.level = level
        else:
            text = item
        run = p.add_run()
        run.text = text
        set_run(run, size - (2 if p.level else 0), False, INK)
        p.space_after = Pt(8 if spacing else 4)
    return box


def add_card(slide, left, top, width, height, title, body_lines, accent=BRAND):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = SOFT
    # accent strip
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.12), height)
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent
    strip.line.fill.background()

    tb = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2), width - Inches(0.4), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, 15, True, BRAND)

    body = slide.shapes.add_textbox(
        left + Inches(0.3), top + Inches(0.65), width - Inches(0.45), height - Inches(0.8)
    )
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(body_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, 13, False, INK)
        p.space_after = Pt(4)


def section_divider(prs, section_no, title, subtitle, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_bg(slide, BRAND)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.5), Inches(1))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"BAGIAN {section_no}"
    set_run(run, 18, True, SOFT)

    box2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.5), Inches(1.2))
    p = box2.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, 40, True, WHITE)

    box3 = slide.shapes.add_textbox(Inches(0.8), Inches(4.3), Inches(11.5), Inches(0.6))
    p = box3.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = subtitle
    set_run(run, 18, False, SOFT)

    foot = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(12), Inches(0.3))
    p = foot.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"{page}/{total}"
    set_run(run, 11, False, SOFT)
    p.alignment = PP_ALIGN.RIGHT
    return slide


def content_slide(prs, title, bullets, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_topbar(slide)
    add_title(slide, title)
    add_bullets(slide, bullets)
    add_footer(slide, page, total)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # Pre-count total slides for footer
    # We'll build and then we know approximate count; use dynamic numbering
    slides_meta = []

    def track(fn):
        slides_meta.append(fn)

    # ---- Collect slide builders ----
    builders = []

    def add(builder):
        builders.append(builder)

    # 1 Cover
    def s_cover(prs, page, total):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide, BRAND)
        mark = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.9), Inches(0.9)
        )
        mark.fill.solid()
        mark.fill.fore_color.rgb = WHITE
        mark.line.fill.background()
        mtf = mark.text_frame
        mtf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = mtf.paragraphs[0].add_run()
        run.text = "CR"
        set_run(run, 28, True, BRAND)

        t = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11), Inches(1))
        run = t.text_frame.paragraphs[0].add_run()
        run.text = "Facepay Cilegon Raudhatul Jannah"
        set_run(run, 44, True, WHITE)

        s = slide.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(11), Inches(0.8))
        run = s.text_frame.paragraphs[0].add_run()
        run.text = "Dokumentasi Lengkap Sistem\nConfig Demo · Tutorial Penggunaan · Penjelasan Per Halaman"
        set_run(run, 18, False, SOFT)

        f = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11), Inches(0.4))
        run = f.text_frame.paragraphs[0].add_run()
        run.text = "ICT  ·  Demo Face Recognition & FacePay  ·  2026"
        set_run(run, 14, False, SOFT)

    add(s_cover)

    # 2 Agenda
    def s_agenda(prs, page, total):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide)
        add_topbar(slide)
        add_title(slide, "Agenda")
        items = [
            ("01", "Pengenalan Sistem", "Apa itu Facepay Cilegon Raudhatul Jannah & arsitektur"),
            ("02", "Nilai Config Demo", "Isi admin/api/config.php untuk lingkungan demo"),
            ("03", "Tutorial Penggunaan", "Alur kerja Admin, Kantin, dan Portal Ortu"),
            ("04", "Penjelasan Per Halaman", "Fungsi setiap halaman & menu"),
            ("05", "Referensi Teknis", "API, autentikasi, tips operasional"),
        ]
        for i, (num, title, desc) in enumerate(items):
            y = Inches(1.7) + Inches(i * 0.95)
            num_box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), y, Inches(0.85), Inches(0.75)
            )
            num_box.fill.solid()
            num_box.fill.fore_color.rgb = BRAND
            num_box.line.fill.background()
            np = num_box.text_frame.paragraphs[0]
            np.alignment = PP_ALIGN.CENTER
            run = np.add_run()
            run.text = num
            set_run(run, 18, True, WHITE)

            tb = slide.shapes.add_textbox(Inches(1.8), y + Inches(0.05), Inches(10), Inches(0.35))
            run = tb.text_frame.paragraphs[0].add_run()
            run.text = title
            set_run(run, 20, True, BRAND_DEEP)

            db = slide.shapes.add_textbox(Inches(1.8), y + Inches(0.4), Inches(10), Inches(0.3))
            run = db.text_frame.paragraphs[0].add_run()
            run.text = desc
            set_run(run, 14, False, MUTED)
        add_footer(slide, page, total)

    add(s_agenda)

    # 3 What is
    def s_intro(prs, page, total):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide)
        add_topbar(slide)
        add_title(slide, "Apa itu Facepay Cilegon Raudhatul Jannah?")
        add_bullets(
            slide,
            [
                "Platform face recognition terpadu untuk sekolah / lembaga pendidikan.",
                "Menggabungkan 3 layanan utama:",
                (1, "FacePay Kantin — pembayaran belanja siswa dengan pemindaian wajah"),
                (1, "Panel Admin — kelola data siswa, rekam wajah, pantau saldo, uji pengenalan"),
                (1, "Portal Orang Tua — rekam / unggah foto referensi wajah anak secara mandiri"),
                "Teknologi: HTML5, CSS3, JavaScript, PHP Proxy API, MySQL, face-api.js, Merchant API VPS.",
            ],
        )
        add_footer(slide, page, total)

    add(s_intro)

    # 4 Architecture
    def s_arch(prs, page, total):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide)
        add_topbar(slide)
        add_title(slide, "Arsitektur Akses (4 Jalur)")
        cards = [
            ("Portal Utama /", ["Pintu masuk publik", "Pilih: Kantin / Admin / Ortu"], BRAND),
            ("Kantin /kantin/", ["Login petugas (Merchant JWT)", "Terminal FacePay + log"], ACCENT),
            ("Admin /admin/", ["Login DB lokal", "Data siswa, rekam, saldo, uji wajah"], BRAND),
            ("Ortu /ortu/", ["Login NIS siswa", "Rekam / unggah foto wajah"], BRAND_DEEP),
        ]
        for i, (title, lines, color) in enumerate(cards):
            x = Inches(0.5) + Inches(i * 3.2)
            add_card(slide, x, Inches(1.9), Inches(3.0), Inches(3.8), title, lines, color)
        add_footer(slide, page, total)

    add(s_arch)

    # 5 Roles
    def s_roles(prs, page, total):
        content_slide(
            prs,
            "Peran & Autentikasi",
            [
                "Guest — tanpa login, hanya portal utama (/).",
                "Petugas Kantin — autentikasi Merchant API (LoginRequest) → sesi PHP CILEGON_RAUDHATUL_JANNAH_KANTIN.",
                "Administrator — tabel admin_user (DB lokal) → sesi PHP MAF_ADMIN.",
                "  Akun default pertama kali: admin / admin123 (wajib diganti).",
                "Orang Tua — verifikasi NIS via ortu-siswa.php → sessionStorage ortu_siswa_session.",
            ],
            page,
            total,
        )

    add(s_roles)

    # SECTION 2 — nilai config demo saja (tanpa panduan setup lengkap)
    def s_sec2(prs, page, total):
        section_divider(
            prs,
            "02",
            "Nilai Config Demo",
            "File: admin/api/config.php",
            page,
            total,
        )

    add(s_sec2)

    def s_config_demo(prs, page, total):
        content_slide(
            prs,
            "Nilai Config Demo (admin/api/config.php)",
            [
                "db.host = localhost",
                "db.name = demo_face",
                "db.user = demo_face",
                "db.pass = demo_face",
                "db.auto_migrate = true",
                "api_url = http://103.23.103.43/MobileMerchant/Malang_Alizzah_ForVPS/index.php",
                "jwt_secret = 4ecfd4c24aee85b4b485f9d828aa1b7d",
                "log_username_default = WS_TESTING",
                "cors_origin = *",
                "",
                "Admin login default: admin / admin123",
            ],
            page,
            total,
        )

    add(s_config_demo)

    # SECTION 3 Tutorial
    def s_sec3(prs, page, total):
        section_divider(
            prs, "03", "Tutorial Penggunaan", "Alur kerja Admin · Kantin · Orang Tua", page, total
        )

    add(s_sec3)

    def s_flow_admin(prs, page, total):
        content_slide(
            prs,
            "Tutorial Admin — Alur Harian",
            [
                "1. Buka portal → Masuk Admin → login (admin / password).",
                "2. Data siswa: sync dari Merchant (StudentRequest) atau input/impor manual.",
                "3. Rekam data: pilih siswa → aktifkan kamera → simpan foto referensi.",
                "4. Saldo siswa: pantau / refresh inquiry saldo ke VPS.",
                "5. Uji wajah: pastikan pengenalan wajah bekerja sebelum dipakai kantin.",
                "6. (Opsional) Setting Unit / Jenjang / Kelas untuk struktur sekolah.",
            ],
            page,
            total,
        )

    add(s_flow_admin)

    def s_flow_siswa(prs, page, total):
        content_slide(
            prs,
            "Tutorial — Kelola Data Siswa",
            [
                "Buka Admin → Data siswa.",
                "Sinkron: tombol Sinkron menarik data dari Merchant API (StudentRequest).",
                "Tambah / Edit / Hapus siswa lewat form modal.",
                "Impor Excel (.xlsx/.xls) atau Ekspor untuk backup.",
                "Filter berdasarkan unit / jenjang / kelas bila master data sudah diisi.",
                "Pastikan NIS unik — dipakai login ortu & inquiry saldo.",
            ],
            page,
            total,
        )

    add(s_flow_siswa)

    def s_flow_rekam(prs, page, total):
        content_slide(
            prs,
            "Tutorial — Rekam Foto Wajah (Admin)",
            [
                "Buka Admin → Rekam data.",
                "Pilih siswa dari daftar (atau cari NIS/nama).",
                "Aktifkan kamera → posisikan wajah di frame → ambil foto.",
                "Bisa ganti kamera depan/belakang; atau unggah dari galeri.",
                "Simpan — foto tersimpan di database sebagai referensi face recognition.",
                "Status: siswa dengan foto siap dikenali di kantin & uji wajah.",
            ],
            page,
            total,
        )

    add(s_flow_rekam)

    def s_flow_kantin(prs, page, total):
        content_slide(
            prs,
            "Tutorial Kantin — FacePay",
            [
                "1. Portal → Masuk Kantin → login username/password petugas (Merchant).",
                "2. Aktifkan kamera pada terminal FacePay.",
                "3. Siswa menghadap kamera → sistem mengenali wajah → tampil nama & foto.",
                "4. Sistem otomatis InquirySALDO (cek sisa saldo).",
                "5. Kasir input nominal belanja → tekan Bayar (PaymentBELANJAKantin).",
                "6. Transaksi sukses tercatat di Log Pembelian kantin.",
                "Tips: tekan Muat ulang wajah jika data referensi baru saja diubah.",
            ],
            page,
            total,
        )

    add(s_flow_kantin)

    def s_flow_ortu(prs, page, total):
        content_slide(
            prs,
            "Tutorial Portal Orang Tua",
            [
                "1. Dari portal utama → Buka portal ortu (atau /ortu/).",
                "2. Masukkan NIS anak → sistem verifikasi ke database.",
                "3. Dashboard menampilkan identitas & status foto (sudah ada / belum).",
                "4. Rekam wajah: aktifkan kamera, ikuti oval guide, ambil foto, simpan.",
                "5. Alternatif: unggah foto dari galeri HP/laptop.",
                "6. Setelah tersimpan, foto dipakai FacePay kantin & uji wajah admin.",
            ],
            page,
            total,
        )

    add(s_flow_ortu)

    # SECTION 4 Pages
    def s_sec4(prs, page, total):
        section_divider(
            prs, "04", "Penjelasan Per Halaman", "Fungsi setiap URL / menu di sistem", page, total
        )

    add(s_sec4)

    def s_pages_portal(prs, page, total):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide)
        add_topbar(slide)
        add_title(slide, "Portal & Autentikasi")
        add_card(
            slide, Inches(0.5), Inches(1.7), Inches(4.0), Inches(4.5),
            "index.html — Portal Utama",
            [
                "Landing page brand Facepay Cilegon Raudhatul Jannah",
                "3 pintu: Kantin, Admin, Ortu",
                "Tanpa login",
            ],
            BRAND,
        )
        add_card(
            slide, Inches(4.7), Inches(1.7), Inches(4.0), Inches(4.5),
            "admin/login.html",
            [
                "Form login administrator",
                "Auth: DB admin_user",
                "Default: admin / admin123",
                "Ganti password setelah masuk",
            ],
            BRAND,
        )
        add_card(
            slide, Inches(8.9), Inches(1.7), Inches(4.0), Inches(4.5),
            "kantin/login.html",
            [
                "Form login petugas kantin",
                "Auth: Merchant LoginRequest",
                "Membuat sesi kantin",
            ],
            ACCENT,
        )
        add_footer(slide, page, total)

    add(s_pages_portal)

    def s_pages_admin_home(prs, page, total):
        content_slide(
            prs,
            "Admin — Beranda (admin/index.html)",
            [
                "Dashboard panel admin setelah login.",
                "Tile pintas ke: Data siswa, Rekam data, Saldo siswa, Uji wajah.",
                "Bottom navigation untuk berpindah modul dengan cepat.",
                "Menampilkan brand Facepay Cilegon Raudhatul Jannah di top bar.",
            ],
            page,
            total,
        )

    add(s_pages_admin_home)

    def s_pages_data(prs, page, total):
        content_slide(
            prs,
            "Admin — Data Siswa (settings/data-siswa.html)",
            [
                "Tabel master siswa (NIS, nama, unit, jenjang, kelas, dll).",
                "CRUD: tambah, edit, hapus.",
                "Sinkron dari Merchant API (StudentRequest).",
                "Impor / ekspor Excel.",
                "Filter & pagination untuk data besar.",
                "API: admin/api/siswa-db.php, siswa-sync.php",
            ],
            page,
            total,
        )

    add(s_pages_data)

    def s_pages_rekam(prs, page, total):
        content_slide(
            prs,
            "Admin — Rekam Data (settings/rekam-data.html)",
            [
                "Pilih siswa → rekam foto wajah referensi via webcam.",
                "Pratinjau foto yang sudah tersimpan; opsi hapus foto.",
                "Ganti kamera depan/belakang; unggah file sebagai alternatif.",
                "Status foto per siswa (sudah / belum).",
                "API: rekam-data.php, rekam-simpan.php, rekam-hapus-foto.php",
            ],
            page,
            total,
        )

    add(s_pages_rekam)

    def s_pages_saldo(prs, page, total):
        content_slide(
            prs,
            "Admin — Saldo Siswa (settings/saldo-siswa.html)",
            [
                "Daftar siswa dengan info saldo (inquiry ke VPS).",
                "Refresh saldo per siswa.",
                "Tautan ke cek saldo berbasis wajah (modules/cek-saldo.html).",
                "API: saldo.php, saldo-inquiry.php",
            ],
            page,
            total,
        )

    add(s_pages_saldo)

    def s_pages_settings(prs, page, total):
        content_slide(
            prs,
            "Admin — Setting Master (Unit / Jenjang / Kelas)",
            [
                "settings/unit.html — master unit sekolah (mis. Putra / Putri).",
                "settings/jenjang.html — tingkat/jenjang (terhubung unit).",
                "settings/kelas.html — kelas (terhubung jenjang).",
                "Digunakan untuk filter & pengelompokan data siswa.",
                "Bisa diisi manual atau diselaraskan dari data SIE/Merchant.",
            ],
            page,
            total,
        )

    add(s_pages_settings)

    def s_pages_modules(prs, page, total):
        content_slide(
            prs,
            "Admin — Modul Uji & FacePay",
            [
                "modules/face.html — Uji pengenalan wajah (deteksi + match referensi).",
                "modules/cek-saldo.html — Cek saldo dengan scan wajah (inquiry).",
                "modules/facepay.html — Modul facepay di sisi admin (uji alur bayar).",
                "Berguna untuk training petugas & validasi sebelum operasional kantin.",
            ],
            page,
            total,
        )

    add(s_pages_modules)

    def s_pages_kantin(prs, page, total):
        content_slide(
            prs,
            "Kantin — Terminal FacePay (kantin/index.html)",
            [
                "Area kamera fullscreen + status deteksi wajah.",
                "Tombol: Aktifkan kamera, Hentikan, Ganti kamera, Muat ulang wajah.",
                "Modal konfirmasi siswa + input nominal + tombol Bayar.",
                "Panel log transaksi / riwayat pembelian sesi berjalan.",
                "API lokal: kantin/api/login.php, saldo-inquiry.php, payment.php, log-transaksi.php",
            ],
            page,
            total,
        )

    add(s_pages_kantin)

    def s_pages_ortu(prs, page, total):
        content_slide(
            prs,
            "Portal Ortu — Halaman",
            [
                "ortu/index.html — Login dengan NIS siswa.",
                "ortu/dashboard.html — Identitas siswa + status foto + form rekam.",
                "Oval face guide membantu posisi wajah yang baik.",
                "Opsi unggah dari galeri jika webcam tidak tersedia.",
                "API: admin/api/ortu-siswa.php + endpoint rekam yang sama dengan admin.",
            ],
            page,
            total,
        )

    add(s_pages_ortu)

    # SECTION 5 Technical
    def s_sec5(prs, page, total):
        section_divider(
            prs, "05", "Referensi Teknis", "API Merchant, proxy PHP, tips operasional", page, total
        )

    add(s_sec5)

    def s_api_merchant(prs, page, total):
        content_slide(
            prs,
            "Merchant API (VPS) — Method",
            [
                "LoginRequest — autentikasi petugas kantin → JWT",
                "InquirySALDO — cek sisa saldo siswa",
                "PaymentBELANJAKantin — potong saldo untuk belanja",
                "LogTransaksiRequest — riwayat transaksi",
                "StudentRequest — tarik / sync data siswa",
                "Semua dipanggil lewat PHP proxy (bukan langsung dari browser ke VPS).",
            ],
            page,
            total,
        )

    add(s_api_merchant)

    def s_api_local(prs, page, total):
        content_slide(
            prs,
            "PHP Proxy API Lokal (cuplikan)",
            [
                "admin/api/auth-login.php / auth-logout.php / auth-me.php / auth-change-password.php",
                "admin/api/siswa-db.php · siswa-sync.php",
                "admin/api/rekam-data.php · rekam-simpan.php · rekam-hapus-foto.php",
                "admin/api/saldo.php · saldo-inquiry.php · health.php",
                "kantin/api/login.php · payment.php · saldo-inquiry.php · log-transaksi.php",
                "admin/api/ortu-siswa.php — verifikasi NIS portal ortu",
            ],
            page,
            total,
        )

    add(s_api_local)

    def s_tips(prs, page, total):
        content_slide(
            prs,
            "Tips Operasional & Troubleshooting",
            [
                "Kamera tidak muncul → pastikan HTTPS / izinkan permission browser.",
                "Wajah tidak dikenali → pastikan foto referensi jelas, cahaya cukup, muat ulang wajah.",
                "Koneksi DB gagal → cek config.php & buat user MySQL dengan hak yang benar.",
                "Login kantin gagal → cek api_url VPS & kredensial merchant.",
                "Saldo tidak update → cek InquirySALDO / jaringan ke VPS.",
                "Ganti password admin segera setelah instalasi pertama.",
            ],
            page,
            total,
        )

    add(s_tips)

    def s_map(prs, page, total):
        content_slide(
            prs,
            "Peta Folder Proyek",
            [
                "/                 → portal utama (index.html)",
                "/admin/           → panel admin (HTML + js + css)",
                "/admin/api/       → PHP backend + config.php (PUSAT KONFIG)",
                "/admin/settings/  → data siswa, rekam, saldo, unit, jenjang, kelas",
                "/admin/modules/   → uji wajah, cek saldo, facepay",
                "/kantin/          → login + terminal FacePay + api kantin",
                "/ortu/            → portal orang tua",
                "/docs/            → dokumentasi & presentasi ini",
            ],
            page,
            total,
        )

    add(s_map)

    def s_closing(prs, page, total):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide, BRAND)
        t = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.5), Inches(1))
        run = t.text_frame.paragraphs[0].add_run()
        run.text = "Siap untuk Demo"
        set_run(run, 40, True, WHITE)

        s = slide.shapes.add_textbox(Inches(0.8), Inches(3.6), Inches(11.5), Inches(1.5))
        tf = s.text_frame
        tf.word_wrap = True
        lines = [
            "Isi config demo → sync siswa → rekam wajah → uji FacePay kantin.",
            "File ini: docs/Facepay_CILEGON_RAUDHATUL_JANNAH_Dokumentasi_Sistem_v2.pptx",
            "Dokumen teks: docs/DOKUMENTASI-SISTEM.md",
        ]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = line
            set_run(run, 16, False, SOFT)
            p.space_after = Pt(8)

        f = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11), Inches(0.4))
        run = f.text_frame.paragraphs[0].add_run()
        run.text = f"Facepay Cilegon Raudhatul Jannah  ·  {page}/{total}"
        set_run(run, 12, False, SOFT)

    add(s_closing)

    total = len(builders)
    for i, builder in enumerate(builders, start=1):
        builder(prs, i, total)

    out = r"d:\PROJECT\ICT\cilegon_raudhatul_jannah\docs\Facepay_CILEGON_RAUDHATUL_JANNAH_Dokumentasi_Sistem_v2.pptx"
    prs.save(out)
    print(f"Saved: {out}")
    print(f"Slides: {total}")
    return out


if __name__ == "__main__":
    build()
